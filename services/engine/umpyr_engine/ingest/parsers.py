"""Parsers for uploaded files.

Tickets: CSV, XLSX, JSON, JSONL (Zendesk/Gorgias-style exports), .eml and .mbox.
Policies: PDF (PyMuPDF), DOCX (python-docx), Markdown/TXT/HTML, or a URL (same-domain crawl, depth <= 2).
Orders: CSV / XLSX / JSON rows keyed by order_id (or booking/PNR).
"""

from __future__ import annotations

import email
import email.policy
import io
import logging
import mailbox
import re
import tempfile
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlparse

from .loader import parse_rows

log = logging.getLogger(__name__)

TICKET_EXTS = {".csv", ".xlsx", ".xlsm", ".json", ".jsonl", ".eml", ".mbox"}
POLICY_EXTS = {".pdf", ".docx", ".md", ".markdown", ".txt", ".html", ".htm"}
RECORD_EXTS = {".csv", ".xlsx", ".xlsm", ".json", ".jsonl"}


# ---------------------------------------------------------------- tickets
def _msg_text(msg) -> str:
    body = msg.get_body(preferencelist=("plain", "html"))
    if body is None:
        return ""
    content = body.get_content()
    if body.get_content_type() == "text/html":
        content = re.sub(r"<[^>]+>", " ", content)
    # drop quoted replies
    lines = []
    for line in content.splitlines():
        if line.strip().startswith(">") or re.match(r"^On .* wrote:$", line.strip()):
            break
        lines.append(line)
    return "\n".join(lines).strip()


def _subject_key(subject: str) -> str:
    return re.sub(r"^(re|fw|fwd)\s*:\s*", "", (subject or "").strip(), flags=re.I).lower()


def parse_mail(data: bytes, suffix: str, staff_domains: list[str] | None = None) -> list[dict[str, Any]]:
    """One ticket per subject thread; messages from staff_domains are role=staff."""
    msgs = []
    if suffix == ".eml":
        msgs = [email.message_from_bytes(data, policy=email.policy.default)]
    else:
        with tempfile.NamedTemporaryFile(suffix=".mbox", delete=False) as f:
            f.write(data)
            path = f.name
        box = mailbox.mbox(path, factory=lambda fp: email.message_from_binary_file(fp, policy=email.policy.default))
        msgs = list(box)
    threads: dict[str, list] = defaultdict(list)
    for msg in msgs:
        threads[_subject_key(msg.get("subject", ""))].append(msg)
    out = []
    for i, ms in enumerate(threads.values(), start=1):
        ms.sort(key=lambda m_: email.utils.parsedate_to_datetime(m_["date"]) if m_.get("date") else datetime.min)
        thread = []
        for m_ in ms:
            sender = (m_.get("from") or "").lower()
            role = "staff" if staff_domains and any(d.lower() in sender for d in staff_domains) else "customer"
            at = None
            try:
                at = email.utils.parsedate_to_datetime(m_["date"]).isoformat() if m_.get("date") else None
            except (TypeError, ValueError):
                pass
            thread.append({"role": role, "at": at, "text": _msg_text(m_)})
        if thread and all(t["role"] == "customer" for t in thread) and len(thread) > 1:
            # heuristic: alternate roles when no staff domain is known
            for j, t in enumerate(thread):
                t["role"] = "customer" if j % 2 == 0 else "staff"
        out.append(
            {
                "ticket_id": f"MAIL-{i:05d}",
                "created_at": thread[0]["at"] if thread else None,
                "channel": "email",
                "subject": ms[0].get("subject", ""),
                "thread": thread,
                "agent_actions": [],
                "resolution": None,
                "amount": None,
                "order_id": None,
            }
        )
    return out


def parse_ticket_file(filename: str, data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """Returns (rows, columns). Mail files are already in the canonical layout."""
    suffix = Path(filename).suffix.lower()
    if suffix in (".eml", ".mbox"):
        rows = parse_mail(data, suffix)
        return rows, [
            "ticket_id",
            "created_at",
            "channel",
            "subject",
            "thread",
            "agent_actions",
            "resolution",
            "amount",
            "order_id",
        ]
    rows = parse_rows(Path(filename), data)
    cols: list[str] = []
    for r in rows[:50]:
        for k in r:
            if k not in cols:
                cols.append(str(k))
    return rows, cols


# ---------------------------------------------------------------- policies
@dataclass
class PolicyText:
    title: str
    text: str
    source_url: str | None = None
    source_type: str = "file"
    notes: list[str] = field(default_factory=list)


def parse_policy_file(filename: str, data: bytes) -> PolicyText:
    suffix = Path(filename).suffix.lower()
    name = Path(filename).stem
    if suffix == ".pdf":
        import pymupdf

        doc = pymupdf.open(stream=data, filetype="pdf")
        text = "\n".join(page.get_text() for page in doc)
        return PolicyText(
            title=(doc.metadata or {}).get("title") or name,
            text=text,
            source_type="pdf",
            notes=[f"pdf {len(doc)} pages"],
        )
    if suffix == ".docx":
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return PolicyText(title=d.core_properties.title or name, text="\n".join(parts), source_type="docx")
    if suffix in (".html", ".htm"):
        import trafilatura

        text = trafilatura.extract(data.decode("utf-8", errors="ignore"), include_tables=True, favor_recall=True) or ""
        return PolicyText(title=name, text=text, source_type="html")
    return PolicyText(
        title=name,
        text=data.decode("utf-8", errors="ignore"),
        source_type="md" if suffix in (".md", ".markdown") else "txt",
    )


def crawl_policy_urls(start_url: str, max_depth: int = 2, max_pages: int = 15) -> list[PolicyText]:
    """Same-domain crawl using the M2 fetcher (handles PDF, Next.js payloads, eCFR, trafilatura)."""
    from ..compile.sources import fetch_source

    domain = urlparse(start_url).netloc
    seen: set[str] = set()
    queue: list[tuple[str, int]] = [(start_url, 0)]
    out: list[PolicyText] = []
    while queue and len(out) < max_pages:
        url, depth = queue.pop(0)
        if url in seen:
            continue
        seen.add(url)
        try:
            doc = fetch_source(f"crawl-{len(seen)}", url)
        except Exception as e:  # noqa: BLE001
            log.warning("crawl fetch failed %s: %s", url, e)
            continue
        if doc.ok:
            out.append(
                PolicyText(
                    title=doc.title or url, text=doc.text, source_url=doc.final_url, source_type="url", notes=doc.notes
                )
            )
        if depth < max_depth and doc.content_type.startswith("text/html"):
            html = doc.raw.decode("utf-8", errors="ignore")
            for href in re.findall(r'href="([^"#]+)"', html):
                link = urljoin(url, href)
                if (
                    urlparse(link).netloc == domain
                    and link not in seen
                    and re.search(
                        r"polic|return|refund|terms|guarantee|warrant|shipping|faq|contract|baggage|fee", link, re.I
                    )
                ):
                    queue.append((link, depth + 1))
    return out


# ---------------------------------------------------------------- records
def parse_record_file(filename: str, data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    rows = parse_rows(Path(filename), data)
    cols: list[str] = []
    for r in rows[:50]:
        for k in r:
            if k not in cols:
                cols.append(str(k))
    return rows, cols
